import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from app.services.report_export import export_business_report


class ReportExportTests(unittest.TestCase):
    def test_excel_report_contains_expected_sheets(self):
        summary = {
            "customers": 1, "offers": 2, "invoices": 3, "revenue": 100.0,
            "open_amount": 50.0, "overdue_amount": 20.0, "overdue_count": 1,
        }
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            export_business_report(
                path, summary, [("2026-09", 100.0)],
                [("R-1", "Kupec", "2026-09-01", "2026-09-15", 50, 0, 50, "Zapadel")],
                [("Kupec", 1, 50, 0)],
            )
            workbook = load_workbook(path, data_only=True)
            self.assertEqual(workbook.sheetnames, [
                "Pregled", "Mesečni prihodki", "Odprte terjatve",
                "Najboljše stranke",
            ])
            self.assertEqual(workbook["Mesečni prihodki"]["B2"].value, 100)
