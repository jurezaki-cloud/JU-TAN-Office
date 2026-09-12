from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.constants import EXPORT_DIR
from app.excel.excel_mapping import MODULES
from app.pdf.pdf_company import load_company
from app.modules.settings.settings_controller import SettingsController
from app.theme.colors import LightColors


HEADER_FILL = PatternFill("solid", fgColor=LightColors.TEXT.replace("#", ""))
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
META_FONT = Font(color=LightColors.SECONDARY.replace("#", ""), name="Calibri", size=10)
TITLE_FONT = Font(color=LightColors.TEXT.replace("#", ""), bold=True, name="Calibri", size=16)
THIN = Border(
    left=Side(style="thin", color=LightColors.BORDER.replace("#", "")),
    right=Side(style="thin", color=LightColors.BORDER.replace("#", "")),
    top=Side(style="thin", color=LightColors.BORDER.replace("#", "")),
    bottom=Side(style="thin", color=LightColors.BORDER.replace("#", "")),
)
ZEBRA = PatternFill("solid", fgColor=LightColors.TABLE_HEADER.replace("#", ""))


def excel_folders() -> dict:
    extras = SettingsController().load_extras()
    excel = extras.get("excel") or {}
    pdf = extras.get("pdf") or {}
    export_folder = excel.get("export_folder") or pdf.get("folder") or str(EXPORT_DIR)
    import_folder = excel.get("import_folder") or export_folder
    Path(export_folder).mkdir(parents=True, exist_ok=True)
    Path(import_folder).mkdir(parents=True, exist_ok=True)
    return {"export": Path(export_folder), "import": Path(import_folder)}


def _maybe_logo(worksheet, logo_path: str) -> int:
    if not logo_path or not Path(logo_path).exists():
        return 1
    try:
        from openpyxl.drawing.image import Image as XLImage
        image = XLImage(logo_path)
        image.width = 96
        image.height = 48
        worksheet.add_image(image, "A1")
        worksheet.row_dimensions[1].height = 42
        return 1
    except Exception:
        return 1


def write_workbook(path: Path, module_key: str, headers: list[str], rows: list[list], template: bool = False) -> Path:
    if len(rows) >= 400:
        return write_workbook_stream(path, module_key, headers, rows)
    company = load_company()
    book = Workbook()
    sheet = book.active
    sheet.title = MODULES[module_key]["title"][:31]
    start = _maybe_logo(sheet, company.logo if not template else "")
    sheet.merge_cells(start_row=1, start_column=2, end_row=1, end_column=max(3, len(headers)))
    sheet.cell(1, 2, company.name or "JU-TAN Office").font = TITLE_FONT
    sheet.cell(2, 2, MODULES[module_key]["title"]).font = META_FONT
    sheet.cell(3, 2, f"Datum izvoza: {date.today().isoformat()}").font = META_FONT
    header_row = 5
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN
    for index, row in enumerate(rows, start=header_row + 1):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(index, column, value)
            cell.border = THIN
            if index % 2 == 0:
                cell.fill = ZEBRA
    last_row = header_row + max(len(rows), 1)
    last_col = get_column_letter(max(len(headers), 1))
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{last_col}{last_row}"
    for column, title in enumerate(headers, start=1):
        width = max(len(str(title)), 12)
        for row in rows:
            if column - 1 < len(row):
                width = max(width, min(len(str(row[column - 1] or "")), 40))
        sheet.column_dimensions[get_column_letter(column)].width = width + 2
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def write_workbook_stream(path: Path, module_key: str, headers: list[str], rows: list[list]) -> Path:
    """Write-only pretok za velike izvoze (manj pomnilnika)."""
    book = Workbook(write_only=True)
    sheet = book.create_sheet(MODULES[module_key]["title"][:31])
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def analyze_file(path: Path) -> dict:
    book = load_workbook(path, data_only=True)
    sheet = book.active
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value).strip() if value is not None else "" for value in (rows[0] if rows else [])]
    # skip branded exports: find header row containing known labels
    header_index = 0
    for index, row in enumerate(rows[:8]):
        values = [str(cell or "").strip() for cell in row]
        joined = " ".join(values).lower()
        if "naziv" in joined or "šifra" in joined or "sifra" in joined or "customer" in joined or "stevilka" in joined or "številka" in joined:
            headers = values
            header_index = index
            break
    data = []
    for row in rows[header_index + 1:]:
        if row is None or all(cell in (None, "") for cell in row):
            continue
        data.append(list(row))
    return {
        "sheet": sheet.title,
        "headers": headers,
        "rows": data,
        "count": len(data),
        "header_row": header_index + 1,
    }


def map_rows(analysis: dict, mapping: dict[str, str]) -> list[dict]:
    headers = analysis["headers"]
    index = {header: i for i, header in enumerate(headers)}
    mapped = []
    for row in analysis["rows"]:
        item = {}
        for key, header in mapping.items():
            if not header or header not in index:
                item[key] = ""
                continue
            value = row[index[header]] if index[header] < len(row) else ""
            item[key] = "" if value is None else value
        mapped.append(item)
    return mapped
