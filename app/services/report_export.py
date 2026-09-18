from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=18, bold=True, color="0F766E")
MONEY_FORMAT = '#,##0.00 [$€-x-euro2]'


def _format_sheet(sheet, widths=None):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for index, width in enumerate(widths or [], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def export_business_report(path, summary, monthly, receivables, top_customers):
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Pregled"
    overview.append(["JU-TAN Office – poslovno poročilo", "Vrednost"])
    overview["A1"].font = TITLE_FONT
    rows = [
        ("Ustvarjeno", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Stranke", summary["customers"]), ("Ponudbe", summary["offers"]),
        ("Računi", summary["invoices"]),
        ("Prejeta plačila letos", summary["revenue"]),
        ("Odprte terjatve", summary["open_amount"]),
        ("Zapadle terjatve", summary["overdue_amount"]),
        ("Število zapadlih računov", summary["overdue_count"]),
    ]
    for row in rows:
        overview.append(row)
    for row in range(6, 9):
        overview.cell(row=row, column=2).number_format = MONEY_FORMAT
    overview.column_dimensions["A"].width = 32
    overview.column_dimensions["B"].width = 22

    monthly_sheet = workbook.create_sheet("Mesečni prihodki")
    monthly_sheet.append(["Mesec", "Prejeta plačila"])
    for month, amount in monthly:
        monthly_sheet.append([month, amount])
        monthly_sheet.cell(monthly_sheet.max_row, 2).number_format = MONEY_FORMAT
    _format_sheet(monthly_sheet, [15, 22])

    receivables_sheet = workbook.create_sheet("Odprte terjatve")
    receivables_sheet.append([
        "Račun", "Stranka", "Izdano", "Zapade", "Skupaj", "Plačano",
        "Odprto", "Status",
    ])
    for row in receivables:
        receivables_sheet.append(list(row))
        for column in (5, 6, 7):
            receivables_sheet.cell(receivables_sheet.max_row, column).number_format = MONEY_FORMAT
    _format_sheet(receivables_sheet, [20, 30, 14, 14, 16, 16, 16, 18])

    customers_sheet = workbook.create_sheet("Najboljše stranke")
    customers_sheet.append(["Stranka", "Računi", "Fakturirano", "Plačano"])
    for row in top_customers:
        customers_sheet.append(list(row))
        customers_sheet.cell(customers_sheet.max_row, 3).number_format = MONEY_FORMAT
        customers_sheet.cell(customers_sheet.max_row, 4).number_format = MONEY_FORMAT
    _format_sheet(customers_sheet, [32, 12, 20, 20])

    workbook.save(output)
    return output
